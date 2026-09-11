from openpyxl import load_workbook
from datetime import datetime
import os

class OpenFinanceLoader:
    def __init__(self, filename="Open_Finance_Datos_Simulados.xlsx"):
        self.filename = filename
        self.clientes_of = []
        self.load_data()

    def load_data(self):
        """Carga los datos de Open Finance desde el archivo Excel"""
        if not os.path.exists(self.filename):
            print(f"Error: {self.filename} no encontrado")
            return

        try:
            wb = load_workbook(self.filename)
            ws = wb['Open Finance']

            # Saltar encabezado en fila 1
            # Estructura: Numero | Nombre | Email | Saldo | Ingresos | Score | Riesgo | Productos | Estado | Deuda | Mora | Fecha
            for row_idx in range(2, ws.max_row + 1):
                cliente_data = {}

                cliente_data['numero_documento'] = ws.cell(row_idx, 1).value
                cliente_data['nombre'] = ws.cell(row_idx, 2).value
                cliente_data['email'] = ws.cell(row_idx, 3).value
                cliente_data['saldo'] = ws.cell(row_idx, 4).value
                cliente_data['ingresos_mensuales'] = ws.cell(row_idx, 5).value
                cliente_data['score_crediticio'] = ws.cell(row_idx, 6).value
                cliente_data['riesgo'] = ws.cell(row_idx, 7).value
                cliente_data['producto'] = ws.cell(row_idx, 8).value
                cliente_data['estado_cuenta'] = ws.cell(row_idx, 9).value
                cliente_data['deuda_vigente'] = ws.cell(row_idx, 10).value
                cliente_data['dias_en_mora'] = ws.cell(row_idx, 11).value
                cliente_data['fecha_ultima_transaccion'] = ws.cell(row_idx, 12).value

                if cliente_data['numero_documento']:
                    self.clientes_of.append(cliente_data)

            print(f"Cargados {len(self.clientes_of)} registros desde Open Finance")
        except Exception as e:
            print(f"Error al cargar Open Finance: {e}")

    def buscar_por_documento(self, numero_documento):
        """Busca registros de Open Finance por número de documento"""
        registros = []
        for cliente in self.clientes_of:
            if cliente['numero_documento'] and str(cliente['numero_documento']).strip() == str(numero_documento).strip():
                registros.append(cliente)
        return registros

    def obtener_datos_resumidos(self, numero_documento):
        """Obtiene un resumen de datos de Open Finance para un cliente"""
        registros = self.buscar_por_documento(numero_documento)

        if not registros:
            return None

        # Calcular totales (convertir a números)
        def to_num(val):
            try:
                return int(val) if val else 0
            except:
                return 0

        saldo_total = sum([to_num(r.get('saldo')) for r in registros])
        ingresos_total = max([to_num(r.get('ingresos_mensuales')) for r in registros])
        score_promedio = sum([to_num(r.get('score_crediticio')) for r in registros]) / len(registros) if registros else 0
        deuda_total = sum([to_num(r.get('deuda_vigente')) for r in registros])
        tiene_mora = any([to_num(r.get('dias_en_mora')) > 0 for r in registros])

        productos = []
        for reg in registros:
            if reg.get('producto') and reg['producto'] not in productos:
                productos.append(reg['producto'])

        return {
            'numero_documento': numero_documento,
            'nombre': registros[0].get('nombre'),
            'email': registros[0].get('email'),
            'saldo_total': saldo_total,
            'ingresos_mensuales': ingresos_total,
            'score_crediticio': int(score_promedio),
            'riesgo': registros[0].get('riesgo'),
            'deuda_vigente': deuda_total,
            'tiene_mora': tiene_mora,
            'dias_mora': max([to_num(r.get('dias_en_mora')) for r in registros]) if registros else 0,
            'productos': productos,
            'cantidad_productos': len(registros),
            'estado_cuenta': registros[0].get('estado_cuenta'),
            'fecha_ultima_transaccion': registros[0].get('fecha_ultima_transaccion'),
            'registros_detalle': registros
        }

    def obtener_todos(self):
        """Retorna todos los registros"""
        return self.clientes_of

    def cliente_encontrado(self, numero_documento):
        """Verifica si un cliente existe en Open Finance"""
        return len(self.buscar_por_documento(numero_documento)) > 0
