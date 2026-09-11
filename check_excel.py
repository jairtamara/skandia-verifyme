from openpyxl import load_workbook

wb = load_workbook('Open_Finance_Datos_Simulados.xlsx')
ws = wb['Open Finance']

print(f"Total filas (incluyendo encabezado): {ws.max_row}")
print(f"Total registros de datos: {ws.max_row - 1}")
print("\nPrimeros registros:")
for row in range(2, min(7, ws.max_row + 1)):
    doc = ws.cell(row, 1).value
    nombre = ws.cell(row, 2).value
    print(f"  {doc} - {nombre}")

print("\nUltimos registros:")
for row in range(max(2, ws.max_row - 3), ws.max_row + 1):
    doc = ws.cell(row, 1).value
    nombre = ws.cell(row, 2).value
    print(f"  {doc} - {nombre}")
